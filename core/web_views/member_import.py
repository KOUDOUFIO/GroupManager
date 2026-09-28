"""Import des membres depuis Excel / CSV : envoi, apercu, confirmation."""

from io import BytesIO

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .. import models
from ..services.import_service import TEMPLATE_HEADERS, ImportFileError, MemberImportService

SESSION_KEY = "member_import_rows"


class MemberImportForm(forms.Form):
    """Fichier a importer et groupe par defaut."""
    file = forms.FileField(label=gettext_lazy("Fichier Excel ou CSV"))
    group = forms.ModelChoiceField(
        queryset=models.Group.objects.order_by("name"), required=False,
        label=gettext_lazy("Ajouter les membres au groupe"), empty_label=gettext_lazy("Aucun (ou colonne « Groupe » du fichier)"),
    )


class MemberImportView(LoginRequiredMixin, PermissionRequiredMixin, View):
    """Etape 1 : envoi du fichier. Etape 2 : apercu. Etape 3 : confirmation."""
    permission_required = "core.add_member"
    raise_exception = True
    template_name = "core/member_import.html"

    def get(self, request):
        """Formulaire d'envoi."""
        request.session.pop(SESSION_KEY, None)
        return render(request, self.template_name, {"form": MemberImportForm()})

    def post(self, request):
        """Apercu apres envoi, ou creation apres confirmation."""
        if request.POST.get("action") == "confirm":
            rows = request.session.pop(SESSION_KEY, None)
            if not rows:
                messages.error(request, _("L'aperçu a expiré. Envoyez à nouveau le fichier."))
                return redirect("member_import")
            created = MemberImportService.create(rows)
            messages.success(request, _("%(count)s membre(s) importé(s).") % {"count": created})
            return redirect("member_list")

        form = MemberImportForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(request, self.template_name, {"form": form})
        try:
            rows = MemberImportService.parse(form.cleaned_data["file"], form.cleaned_data["group"])
        except ImportFileError as error:
            form.add_error("file", str(error))
            return render(request, self.template_name, {"form": form})

        request.session[SESSION_KEY] = MemberImportService.serialize(rows)
        counts = {status: sum(1 for row in rows if row.status == status) for status in ("new", "duplicate", "error")}
        return render(request, self.template_name, {"preview": rows, "counts": counts, "form": form})


@login_required
@permission_required("core.add_member", raise_exception=True)
def member_import_template(request):
    """Modele Excel a remplir, avec une ligne d'exemple."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _("Membres")
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(["Ama Mensah", "90 12 34 56", "ama@example.com", "Lomé, Bè", ""])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B3D1F")
    for column, width in zip("ABCDE", (28, 18, 28, 24, 26)):
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="modele-membres-kotiza.xlsx"'
    return response
