"""Import de membres depuis un fichier Excel (.xlsx) ou CSV.

Deux temps : `parse` lit le fichier et classe chaque ligne (a creer, doublon,
erreur) sans rien enregistrer ; `create` enregistre les lignes validees.
"""

import csv
import io
import re
import unicodedata
from dataclasses import asdict, dataclass

from django.db import transaction
from django.utils.translation import gettext as _

from ..models import Group, Member
from .phone import normalize_phone

MAX_ROWS = 2000
MAX_BYTES = 2 * 1024 * 1024

# En-tetes acceptes (sans accents, en minuscules) pour chaque champ.
HEADERS = {
    "full_name": {"nom complet", "nom", "nom et prenom", "nom et prenoms", "nom prenom", "name", "full name", "membre"},
    "phone": {"telephone", "tel", "phone", "numero", "contact", "whatsapp"},
    "email": {"email", "e-mail", "courriel", "mail"},
    "address": {"adresse", "address", "quartier", "ville"},
    "group": {"groupe", "group", "tontine", "association"},
}
TEMPLATE_HEADERS = ["Nom complet", "Téléphone", "Email", "Adresse", "Groupe"]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class ImportFileError(Exception):
    """Fichier illisible, trop gros ou sans colonne "Nom"."""


@dataclass
class ImportRow:
    line: int
    full_name: str
    phone: str = ""
    email: str = ""
    address: str = ""
    group_id: int = None
    group_name: str = ""
    status: str = "new"  # new | duplicate | error
    reason: str = ""


def _key(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text.replace("_", " ")).strip().lower()


def _read_rows(uploaded_file):
    name = (uploaded_file.name or "").lower()
    if uploaded_file.size and uploaded_file.size > MAX_BYTES:
        raise ImportFileError(_("Fichier trop volumineux (2 Mo maximum)."))
    data = uploaded_file.read()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return list(csv.reader(io.StringIO(text), dialect))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        except Exception as exc:
            raise ImportFileError(_("Fichier Excel illisible.")) from exc
        return [["" if cell is None else cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    raise ImportFileError(_("Format non reconnu : envoyez un fichier .xlsx ou .csv."))


class MemberImportService:
    """Lecture et creation des membres importes."""

    @staticmethod
    def parse(uploaded_file, default_group=None) -> list:
        """Classe chaque ligne du fichier, sans rien enregistrer."""
        rows = _read_rows(uploaded_file)
        header_index = None
        columns = {}
        for index, row in enumerate(rows[:10]):
            found = {}
            for col, cell in enumerate(row):
                for field, names in HEADERS.items():
                    if _key(cell) in names and field not in found:
                        found[field] = col
            if "full_name" in found:
                header_index, columns = index, found
                break
        if header_index is None:
            raise ImportFileError(_("Colonne « Nom complet » introuvable. Utilisez le modèle proposé."))

        data_rows = rows[header_index + 1:]
        if len(data_rows) > MAX_ROWS:
            raise ImportFileError(_("Trop de lignes (%(max)s maximum par import).") % {"max": MAX_ROWS})

        groups = {_key(g.name): g for g in Group.objects.all()}
        existing_phones = {normalize_phone(p) for p in Member.objects.exclude(phone="").values_list("phone", flat=True)}
        existing_emails = {e.lower() for e in Member.objects.exclude(email="").values_list("email", flat=True)}
        existing_names = {_key(n) for n in Member.objects.values_list("full_name", flat=True)}
        seen_phones, seen_emails, seen_names = set(), set(), set()

        def cell(row, field):
            col = columns.get(field)
            return str(row[col]).strip() if col is not None and col < len(row) and row[col] not in (None, "") else ""

        result = []
        for offset, row in enumerate(data_rows, start=header_index + 2):
            full_name = re.sub(r"\s+", " ", cell(row, "full_name"))
            raw_phone, email, address = cell(row, "phone"), cell(row, "email").lower(), cell(row, "address")
            if not (full_name or raw_phone or email):
                continue  # ligne vide
            item = ImportRow(line=offset, full_name=full_name, email=email, address=address[:255])

            # Numero lu comme un nombre par Excel : 90123456.0 -> 90123456
            raw_phone = re.sub(r"\.0$", "", raw_phone)
            phone = normalize_phone(raw_phone) if raw_phone else ""
            item.phone = phone or raw_phone

            group_name = cell(row, "group")
            group = groups.get(_key(group_name)) if group_name else default_group
            if group is not None:
                item.group_id, item.group_name = group.pk, group.name

            if not full_name:
                item.status, item.reason = "error", _("Nom manquant.")
            elif len(full_name) > 200:
                item.status, item.reason = "error", _("Nom trop long.")
            elif raw_phone and not phone:
                item.status, item.reason = "error", _("Téléphone invalide.")
            elif email and not EMAIL_RE.match(email):
                item.status, item.reason = "error", _("Email invalide.")
            elif group_name and group is None:
                item.status, item.reason = "error", _("Groupe « %(name)s » inconnu.") % {"name": group_name}
            elif (phone and (phone in existing_phones or phone in seen_phones)) or (
                email and (email in existing_emails or email in seen_emails)
            ):
                item.status, item.reason = "duplicate", _("Même téléphone ou email qu'un membre existant.")
            elif not phone and not email and (_key(full_name) in existing_names or _key(full_name) in seen_names):
                item.status, item.reason = "duplicate", _("Même nom qu'un membre existant.")

            if item.status == "new":
                seen_phones.add(phone)
                seen_emails.add(email)
                seen_names.add(_key(full_name))
            result.append(item)
        return result

    @staticmethod
    @transaction.atomic
    def create(rows: list) -> int:
        """Cree les membres des lignes valides ; renvoie le nombre de membres crees."""
        created = 0
        for row in rows:
            if row["status"] != "new":
                continue
            member = Member.objects.create(
                full_name=row["full_name"], phone=row["phone"], email=row["email"], address=row["address"],
            )
            if row.get("group_id"):
                member.groups.add(row["group_id"])
            created += 1
        return created

    @staticmethod
    def serialize(rows: list) -> list:
        """Lignes au format stockable en session."""
        return [asdict(row) for row in rows]
